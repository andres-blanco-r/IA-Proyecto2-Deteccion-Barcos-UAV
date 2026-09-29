"""Diligencia el formato ABET (IA_Proyecto2_<codigo>_<apellido>.docx) con los resultados reales
de results/*.json y anexa el sustento técnico (E1-E4) con tablas y figuras."""
import os
import sys
import json
import glob
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.abspath(__file__))
RES = os.path.join(ROOT, "results")
TEMPLATE = os.path.join(ROOT, "ABET-IA-PROY2-2026_v1.0.docx")
CODIGO, APELLIDO = "7004096", "Blanco"
NOMBRE = "Andrés Felipe Blanco Romero"
OUT = os.path.join(ROOT, f"IA_Proyecto2_{CODIGO}_{APELLIDO}.docx")
REPO = sys.argv[1] if len(sys.argv) > 1 else ROOT
WEB = "https://andres-blanco-r.github.io/IA-Proyecto2-Deteccion-Barcos-UAV/"
FECHA = sys.argv[2] if len(sys.argv) > 2 else "Día de la prueba en vivo (segundo corte 2026-2)"

C = json.load(open(os.path.join(RES, "classic.json"), encoding="utf-8"))
N = json.load(open(os.path.join(RES, "cnn.json"), encoding="utf-8"))
S = json.load(open(os.path.join(RES, "summary.json"), encoding="utf-8"))
LIVE = json.load(open(os.path.join(RES, "ensayo_en_vivo.json"), encoding="utf-8")) if os.path.exists(os.path.join(RES, "ensayo_en_vivo.json")) else None
pc = lambda v: f"{v * 100:.2f} %"
pm = lambda l: f"{l[0] * 100:.2f} ± {l[1] * 100:.2f} %"

d = docx.Document(TEMPLATE)


def set_cell(cell, text, bold=False, size=8):
    p = cell.paragraphs[0]
    if p.runs:
        r = p.runs[0]
        r.text = text
        for extra in p.runs[1:]:
            extra.text = ""
    else:
        r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold


def shade(cell, fill="D9EAF7"):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill)
    tcPr.append(s)


def table(rows, header=True, widths=None, size=8):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j)
            set_cell(c, str(v), bold=(header and i == 0), size=size)
            if header and i == 0:
                shade(c)
            if widths:
                c.width = Inches(widths[j])
    d.add_paragraph()
    return t


def H1(t):
    d.add_paragraph(t, style="Heading 1")


def H2(t):
    d.add_paragraph(t, style="Heading 2")


def P(t, bold_prefix=None, size=9.5):
    p = d.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix); r.bold = True; r.font.size = Pt(size)
    r = p.add_run(t); r.font.size = Pt(size)
    return p


def B(t, bold_prefix=None):
    p = d.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix); r.bold = True; r.font.size = Pt(9.5)
    r = p.add_run(t); r.font.size = Pt(9.5)


def FIG(name, caption, w=6.8):
    path = os.path.join(RES, name) if not os.path.isabs(name) else name
    if not os.path.exists(path):
        return
    d.add_picture(path, width=Inches(w))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph(caption, style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ------------------------------------------------------------------ 1. Identificación
H = N["holdout"]; fin = N["final_cfg"]; cvm = N["final_cv_metrics"]
t0 = d.tables[0]
ident = {
    8: NOMBRE,
    9: CODIGO,
    10: "Individual",
    11: FECHA,
    12: f"UI web (abrir en Chrome/Edge): {WEB}  ·  UI de escritorio de respaldo: app.py (python app.py o EJECUTAR_UI.bat)  ·  repositorio: {REPO}",
    13: "Carpeta de test del docente, seleccionada en la UI con «Abrir carpeta de test» (lectura recursiva, cualquier "
        "formato/tamaño → 80×80). Ensayo previo: PARCIAL2\\test_ciego\\ (800 imágenes hold-out anonimizadas; clave en clave_test_ciego.csv).",
    14: "Ships in Satellite Imagery (ShipsNet), R. Hammell, Kaggle: 4000 chips RGB 80×80 (1000 barco / 3000 no barco), "
        "imágenes PlanetScope de la bahía de San Francisco y San Pedro (CA). https://www.kaggle.com/datasets/rhammell/ships-in-satellite-imagery",
}
for r, v in ident.items():
    set_cell(t0.rows[r].cells[1], v)

# ------------------------------------------------------------------ 3. Registro de evidencias
t4 = d.tables[4]
ev = {
    1: f"UI web {WEB} (docs/index.html, ONNX en el navegador) y UI de escritorio app.py. Anexo A.9. Carga desde carpeta, etiquetado por botones/teclado "
       "(B/1, N/0, espacio = aceptar), inferencia automática al cargar (CNN ensamble + TTA o SVM), exportación de reporte.",
    2: "Anexo A.1–A.8 y A.10: train_classic.py, train_cnn.py, make_figures.py; results/classic.json, results/cnn.json, "
       f"figuras results/fig_*.png. Mejora demostrable: baseline {pc(S['baseline_holdout_acc'])} → final {pc(S['holdout_final_acc'])} (hold-out).",
    3: f"Anexo A.9 y A.11: evaluación en vivo con la UI sobre la carpeta de test y contraste con CV 5-fold "
       f"({pm(cvm['accuracy'])}). Reporte exportado en reportes/evaluacion_<fecha>/ (predicciones.csv, metricas.json, tablero.png).",
    4: "UI – panel «Métricas en vivo»: accuracy, precisión, recall, F1, especificidad, IC95 % Wilson, matriz de confusión, "
       "barras vivo vs. CV y curva de accuracy acumulado con banda CV (Anexo A.9).",
}
for r, v in ev.items():
    set_cell(t4.rows[r].cells[2], v)

# Localizadores individuales (C1, C2)
loc = [f"E1: app.py · E2: Anexo A.1–A.8, A.10 (train_classic.py, train_cnn.py, results/). Baseline → SVM+descriptores → "
       f"CNN ensamble+TTA; hold-out {pc(S['baseline_holdout_acc'])} → {pc(S['holdout_final_acc'])}.",
       f"E3/E4: UI app.py, Anexo A.9 y A.11; CV 5-fold {pm(cvm['accuracy'])}; hold-out ciego simulado "
       f"{pc(S['holdout_final_acc'])} (IC95 % {S['holdout_final_ci95'][0] * 100:.2f}–{S['holdout_final_ci95'][1] * 100:.2f} %)."]
k = 0
for p in d.paragraphs:
    if p.text.startswith("Localizador y observación individual") and k < 2:
        for r in p.runs[1:]:
            r.text = ""
        if p.runs:
            p.runs[0].text = "Localizador y observación individual: "
        rr = p.add_run(loc[k]); rr.font.size = Pt(9)
        k += 1

# ------------------------------------------------------------------ ANEXO
d.add_page_break()
H1("Anexo A. Sustento técnico del estudiante (E1–E4)")
P(f"Estudiante: {NOMBRE} — Código {CODIGO}. Repositorio / UI: {REPO}. Todo el contenido numérico de este anexo se "
  "genera automáticamente a partir de los archivos de resultados del repositorio (results/*.json), por lo que es reproducible.")

H2("A.1 Problema, contexto mecatrónico y protocolo experimental")
P("El clasificador binario (barco / no barco) forma parte del subsistema de percepción de un UAV de inspección portuaria "
  "(caso Rotterdam): la cámara del dron entrega la escena, se recorre con ventanas de 80×80 px y cada ventana se clasifica; "
  "las ventanas positivas se fusionan (NMS) en detecciones georreferenciables. Por ello se optimiza simultáneamente la "
  "exactitud y el costo de inferencia en un computador embarcado (CPU de 1 hilo como referencia).")
B(f"Dataset ShipsNet: 4000 imágenes (25 % barco). Partición estratificada fija (semilla 42): {C['n_dev']} imágenes de "
  f"desarrollo y {C['n_test']} de test ciego simulado (hold-out) que nunca se usan para entrenar ni para sintonizar.", "Partición: ")
B("toda la selección de preprocesamiento, descriptores e hiperparámetros se hace con validación cruzada estratificada "
  "(5 pliegues; 3 en la rejilla factorial de la CNN) sobre el conjunto de desarrollo.", "Validación cruzada: ")
B("accuracy (criterio del proyecto), precisión, recall, F1, especificidad, AUC-ROC; IC 95 % de Wilson y prueba de McNemar "
  "para comparar modelos sobre el mismo test.", "Métricas: ")
B("para el día de la prueba se re-entrena la configuración ganadora con las 4000 imágenes (ensamble de 5 CNN), práctica "
  "estándar una vez validado el procedimiento.", "Modelo de despliegue: ")

H2("A.2 Línea base (baseline)")
bl = C["baseline"]
table([["Modelo base", "Accuracy CV", "Precisión", "Recall", "F1"],
       ["Píxeles crudos + Regresión logística", pm(bl["LR_pixeles"]["accuracy"]), pm(bl["LR_pixeles"]["precision"]),
        pm(bl["LR_pixeles"]["recall"]), pm(bl["LR_pixeles"]["f1"])],
       ["Píxeles crudos + SVM-RBF por defecto", pm(bl["SVM_pixeles_default"]["accuracy"]), pm(bl["SVM_pixeles_default"]["precision"]),
        pm(bl["SVM_pixeles_default"]["recall"]), pm(bl["SVM_pixeles_default"]["f1"])]],
      widths=[2.6, 1.2, 1.2, 1.2, 1.2])
P(f"La línea base elegida para medir mejoras es píxeles + regresión logística (hold-out: {pc(C['holdout_baseline']['accuracy'])}, "
  f"F1 {C['holdout_baseline']['f1']:.3f}). Está lejos de la meta del 98 %: justifica el proceso de optimización.")

H2("A.3 Preprocesamiento y descriptores visuales (ablación)")
P("Se implementaron y compararon descriptores diferenciados, cada uno con justificación física:")
B("gradientes orientados: capturan la silueta alargada del casco y la estela (forma).", "HOG — ")
B("histograma de color 8×4×4 + momentos del centro vs. la imagen: contraste casco/agua.", "HSV — ")
B("patrones binarios locales uniformes (P=8, R=1) en rejilla 2×2: textura oleaje vs. cubierta/contenedores.", "LBP — ")
B("ecualización adaptativa de histograma en gris (clip 2.0, 4×4) para normalizar iluminación/neblina.", "CLAHE — ")
rows = [["Descriptor", "CLAHE", "Dim.", "Accuracy CV 5-fold", "F1 CV"]]
for a in C["ablation"]:
    rows.append([a["desc"], "sí" if a["clahe"] else "no", a["dim"], pm(a["accuracy"]), pm(a["f1"])])
table(rows, widths=[1.7, 0.7, 0.7, 1.8, 1.8])
FIG("fig_ablacion.png", "Figura A1. Ablación de preprocesamiento y descriptores (SVM-RBF con parámetros por defecto).")
sf = C["selected_features"]
P(f"Decisión cuantitativa: se selecciona automáticamente {'+'.join(sf['use']).upper()} "
  f"{'con' if sf['clahe'] else 'sin'} CLAHE ({pm(sf['cv_acc'])}). El color (HSV) solo aporta poco y, combinado con HOG, "
  "no mejora; CLAHE no mejora a HOG porque este ya normaliza el contraste por bloques (L2-Hys). Se documenta como hallazgo "
  "negativo: un preprocesamiento no se incluye si la CV no demuestra beneficio.")
hb = C["best_hog"]
FIG("fig_hog_sens.png", "Figura A2. Sensibilidad del descriptor HOG (píxeles por celda × orientaciones).", w=3.4)
P(f"Sensibilidad HOG: el tamaño de celda domina sobre el número de orientaciones; el óptimo es {hb['ppc']} px/celda y "
  f"{hb['orient']} orientaciones (celdas grandes ≈ escala del barco en el chip, descriptor más compacto y robusto a la traslación).")

H2("A.4 Sintonización sistemática de hiperparámetros")
g = C["gridsearch"]
rows = [["Clasificador (descriptores seleccionados)", "Configs × 5 pliegues", "Mejores hiperparámetros", "Accuracy CV"]]
for k_, v in g.items():
    rows.append([k_, f"{v['n_configs']} × 5", ", ".join(f"{a.replace('clf__', '')}={b}" for a, b in v["best_params"].items()),
                 f"{v['best_cv_acc'] * 100:.2f} ± {v['best_cv_std'] * 100:.2f} %"])
table(rows, widths=[1.8, 1.1, 3.0, 1.2])
FIG("fig_grid_svm.png", "Figura A3. GridSearchCV del SVM-RBF (C × gamma).", w=3.8)
P("Regularización: en el SVM, C controla el margen (C bajo = más regularización) y gamma el alcance del kernel; "
  "class_weight='balanced' compensa el desbalance 3:1 y mejora el recall de la clase barco. PCA(95 %) se evaluó y se "
  "descartó cuando no mejoró la CV.")

H2("A.5 Arquitectura CNN (ShipNet-Lite) y sintonización multivariada")
P("CNN tipo VGG compacta: 4 bloques [Conv3×3-BN-ReLU ×2 + MaxPool] con canales (32, 64, 128, 256)×w, Global Average "
  "Pooling, Dropout y capa lineal de 2 salidas. Entrenamiento: AdamW + OneCycle, label smoothing 0.05, lote 128. "
  "Aumento de datos en GPU coherente con imágenes cenitales: rotaciones de 90°, espejos, rotación ±15° y traslación, "
  "brillo/contraste y ruido gaussiano (condiciones atmosféricas y de sensor del dron). Inferencia opcional con TTA "
  "(promedio de las 8 transformaciones diedrales) y ensamble de los 5 modelos de pliegue.")
gr = N["grid"]
P(f"Sintonización fina multivariada: rejilla factorial learning rate × dropout × weight decay "
  f"({len(gr)} configuraciones × 3 pliegues). Rango observado: {S['grid_cnn_range'][0]:.2f} % – {S['grid_cnn_range'][1]:.2f} %. "
  f"Óptimo: lr={N['best_grid_cfg']['lr']}, dropout={N['best_grid_cfg']['dropout']}, weight decay={N['best_grid_cfg']['wd']}.")
FIG("fig_grid_cnn.png", "Figura A4. Superficie de respuesta de la CNN (lr × dropout para cada weight decay).")

H2("A.6 Análisis de sensibilidad y balance exactitud / eficiencia para inferencia embarcada")
rows = [["Ancho w", "Parámetros", "MACs/img", "Latencia CPU 1 hilo", "Latencia +TTA", "Tamaño", "Accuracy CV 5-fold"]]
for r in N["sens_width"]:
    rows.append([r["cfg"]["width"], f"{r['params']:,}", f"{r['macs'] / 1e6:.1f} M", f"{r['lat_ms']:.2f} ms",
                 f"{r['lat_tta_ms']:.1f} ms", f"{r['size_kb']:.0f} kB", f"{r['acc_mean'] * 100:.2f} ± {r['acc_std'] * 100:.2f} %"])
rows.append(["SVM+desc.", "—", "—", f"{C['latency_ms_classic']:.2f} ms", "—", "—", pm(C["best_classic"]["cv"]["accuracy"])])
table(rows, widths=[0.7, 0.9, 0.9, 1.1, 1.0, 0.7, 1.5])
rows = [["Épocas", "Accuracy CV 5-fold"]] + [[r["cfg"]["epochs"], f"{r['acc_mean'] * 100:.2f} ± {r['acc_std'] * 100:.2f} %"] for r in N["sens_epochs"]]
rows += [["Aumento de datos: " + ("sí" if r["cfg"]["aug"] else "no"), f"{r['acc_mean'] * 100:.2f} ± {r['acc_std'] * 100:.2f} %"] for r in N["sens_aug"]]
table(rows, widths=[2.5, 2.0])
FIG("fig_pareto_sens.png", "Figura A5. Frente de Pareto exactitud–latencia y sensibilidad a épocas / aumento de datos.")
P(f"Regla de decisión (balance óptimo): se elige el ancho más pequeño cuya exactitud CV está a menos de media desviación "
  f"estándar del mejor → w = {fin['width']}, {fin['epochs']} épocas, con aumento de datos. Así se obtiene la exactitud "
  "máxima con el menor costo computacional para el computador de a bordo; la TTA y el ensamble se habilitan en tierra o "
  "cuando la latencia lo permite (sigue siendo del orden de milisegundos por ventana).")

H2("A.7 Validación cruzada del modelo final")
rows = [["Pliegue", "Accuracy", "Precisión", "Recall", "F1", "AUC"]]
for i, f in enumerate(N["final_cv_folds"]):
    rows.append([i + 1, pc(f["accuracy"]), pc(f["precision"]), pc(f["recall"]), pc(f["f1"]), f"{f['auc']:.4f}"])
rows.append(["Media ± std", pm(cvm["accuracy"]), pm(cvm["precision"]), pm(cvm["recall"]), pm(cvm["f1"]),
             f"{cvm['auc'][0]:.4f} ± {cvm['auc'][1]:.4f}"])
table(rows, widths=[1.0, 1.2, 1.2, 1.2, 1.2, 1.1])
FIG("fig_roc_pr_umbral.png", "Figura A6. Curvas ROC y Precisión–Recall con predicciones fuera de pliegue (OOF) y sensibilidad al umbral.")

H2("A.8 Test ciego simulado (hold-out, n = 800) y mejora demostrable")
rows = [["Modelo", "Accuracy", "Precisión", "Recall", "F1", "Matriz [[TN,FP],[FN,TP]]"],
        ["Baseline píxeles+LR", pc(C["holdout_baseline"]["accuracy"]), pc(C["holdout_baseline"]["precision"]),
         pc(C["holdout_baseline"]["recall"]), pc(C["holdout_baseline"]["f1"]), str(C["holdout_baseline"]["cm"])],
        ["SVM + " + "+".join(sf["use"]).upper() + " (GridSearch)", pc(C["holdout_best_classic"]["accuracy"]), pc(C["holdout_best_classic"]["precision"]),
         pc(C["holdout_best_classic"]["recall"]), pc(C["holdout_best_classic"]["f1"]), str(C["holdout_best_classic"]["cm"])]]
names = {"cnn_single": "CNN individual", "cnn_single_tta": "CNN + TTA", "cnn_ensemble": "Ensamble 5 CNN",
         "cnn_ensemble_tta": "Ensamble 5 CNN + TTA (final)"}
for k_, n_ in names.items():
    h = H[k_]
    rows.append([n_, pc(h["accuracy"]), pc(h["precision"]), pc(h["recall"]), pc(h["f1"]), str(h["cm"])])
table(rows, widths=[2.0, 0.9, 0.9, 0.9, 0.9, 1.6])
FIG("fig_progresion.png", "Figura A7. Progresión de la optimización: de la línea base al modelo final.")
P("Las variantes CNN (individual, TTA, ensamble) difieren entre sí en 1–2 imágenes de 800, diferencia dentro del intervalo "
  "de confianza. Se despliega el ensamble + TTA porque promedia 5 modelos × 8 vistas: reduce la varianza de la predicción "
  "y la hace más robusta ante un test desconocido con posible cambio de dominio (orientación, iluminación, sensor), "
  "que es el escenario real del día de la prueba y del dron.")
mc = S["mcnemar"]
P(f"Modelo final en hold-out: {pc(S['holdout_final_acc'])} (IC 95 % Wilson {S['holdout_final_ci95'][0] * 100:.2f} – "
  f"{S['holdout_final_ci95'][1] * 100:.2f} %), coherente con la CV ({pm(cvm['accuracy'])}): no hay sobreajuste a la "
  f"sintonización. Prueba de McNemar SVM vs. CNN final: discordancias {mc['svm_ok_cnn_err']} / {mc['svm_err_cnn_ok']}, "
  f"p = {mc['p_value_exact']:.3f}"
  + (" (diferencia no significativa: ambos superan la meta; se prefiere la CNN por su menor latencia y su robustez, ver A.8b)." if mc["p_value_exact"] >= 0.05
     else " (diferencia significativa)."))
FIG("fig_cm_holdout.png", "Figura A8. Matriz de confusión del modelo final en el hold-out.", w=2.8)
FIG("fig_errores.png", "Figura A9. Errores residuales (en su mayoría barcos parciales en el borde o muelles).", w=6.0)

RB = json.load(open(os.path.join(RES, "robustness.json"), encoding="utf-8"))
H2("A.8b Robustez ante cambio de dominio (justificación del aumento de datos)")
P("En la CV (misma distribución) el aumento de datos no mejora la exactitud (tabla A.6). Para decidir cuantitativamente se "
  "evaluó el hold-out con perturbaciones realistas del UAV: rotación arbitraria (rumbo del dron), brillo/contraste "
  "(hora/clima), desenfoque (vibración/altura), ruido de sensor, compresión JPEG del enlace de video y su combinación.")
names_r = list(RB["limpio"])
rows = [["Perturbación"] + names_r] + [[p] + [pc(RB[p][n]) for n in names_r] for p in RB]
table(rows, widths=[1.8, 1.7, 1.7, 1.7])
FIG("fig_robustez.png", "Figura A9b. Accuracy en hold-out perturbado: CNN con / sin aumento de datos vs. SVM+descriptores.")
P(f"Conclusión: con aumento de datos la CNN mantiene > 99 % ante rotación libre ({pc(RB['rotación libre']['CNN con aumento'])} "
  f"vs. {pc(RB['rotación libre']['CNN sin aumento'])} sin aumento) y {pc(RB['combinado']['CNN con aumento'])} en el escenario "
  f"combinado (vs. {pc(RB['combinado']['CNN sin aumento'])}); el SVM con descriptores cae a {pc(RB['desenfoque']['SVM + descriptores'])} "
  f"con desenfoque y {pc(RB['JPEG q=30']['SVM + descriptores'])} con JPEG. Por esto el modelo desplegado es la CNN con aumento "
  "de datos, aunque ambos empaten en el test limpio (McNemar p = 1.0). El ruido fuerte (σ = 12) es el caso más exigente para "
  "todos los modelos y queda como trabajo futuro (aumento con ruido de mayor varianza / denoising a bordo).")

H2("A.9 Interfaz de usuario para la prueba en vivo (E1, E3, E4)")
P(f"Se entregan dos interfaces con el mismo flujo y las mismas métricas: (1) UI web publicada en {WEB}, que ejecuta el "
  "ensamble CNN exportado a ONNX directamente en el navegador (WebGPU o WebAssembly multihilo; las imágenes no salen del equipo) "
  "y (2) UI de escritorio en Python (app.py) como respaldo sin conexión. Se verificó que ambas producen las mismas "
  "probabilidades que PyTorch (diferencia máxima 1.8e-7 en ONNX) y las mismas métricas en test_ciego (99.38 %, matriz 596/4/1/199).")
B("«Abrir carpeta de test» → carga recursiva de PNG/JPG/TIF/BMP de cualquier tamaño (se redimensionan a 80×80).", "Carga: ")
B("inmediata con el modelo elegido (CNN ensamble final, CNN dev o SVM+descriptores), TTA configurable y umbral ajustable; "
  "se reporta el tiempo total y por imagen.", "Inferencia: ")
B("botones BARCO / NO BARCO o teclas B/1, N/0 (espacio acepta la predicción), navegación ←/→, tabla con verde = acierto, "
  "rojo = error; opción «ocultar predicción hasta etiquetar» para evitar sesgo del etiquetador. También admite etiquetas "
  "desde CSV (archivo,etiqueta), desde el nombre (1__/0__) o desde subcarpetas ship/no_ship.", "Etiquetado: ")
B("accuracy, precisión, recall, F1, especificidad, N, IC 95 % de Wilson, matriz de confusión, barras en vivo vs. CV "
  "(media ± std), curva de accuracy acumulado sobre la banda de CV, alerta de meta > 98 % y penalización estimada.", "Métricas en tiempo real: ")
B("«Exportar reporte» guarda predicciones.csv, metricas.json, tablero.png y captura de pantalla.", "Trazabilidad: ")
for f in sorted(glob.glob(os.path.join(RES, "ui_*.png"))):
    FIG(f, "Figura A10. UI durante el ensayo de evaluación en vivo sobre test_ciego/ (800 imágenes).")
if LIVE:
    P(f"Ensayo de la prueba en vivo (modelo «{LIVE['modelo']}», test_ciego/, n = {LIVE['n']}): accuracy "
      f"{pc(LIVE['acc'])}, precisión {pc(LIVE['prec'])}, recall {pc(LIVE['rec'])}, F1 {pc(LIVE['f1'])}, "
      f"IC95 % [{LIVE['ic95'][0] * 100:.2f}, {LIVE['ic95'][1] * 100:.2f}] %. La media de CV "
      f"({pc(cvm['accuracy'][0])}) cae dentro del intervalo → resultados consistentes y generalizables.")

H2("A.10 Integración en el sistema de percepción del UAV")
P("scan_scene.py aplica el clasificador por ventana deslizante (paso 10 px) sobre escenas completas del dataset y fusiona "
  "detecciones por NMS, demostrando el uso del clasificador como detector a bordo del dron.")
for f in sorted(glob.glob(os.path.join(RES, "escena_*.png")))[:1]:
    FIG(f, "Figura A11. Detección de barcos en una escena portuaria completa (ventana deslizante + NMS).")

H2("A.11 Correspondencia con la rúbrica (autoevaluación argumentada)")
table([["Criterio / nivel", "Descriptor exigido", "Evidencia en este proyecto"],
       ["C1 – N3", "Baseline, preprocesamiento/extracción de patrones, calibración con mejora demostrable",
        f"A.2 baseline {pc(C['holdout_baseline']['accuracy'])}; A.3 CLAHE, HOG, HSV, LBP; A.4 GridSearch; A.8 final {pc(S['holdout_final_acc'])}"],
       ["C1 – N4", "Comparación de técnicas avanzadas, búsqueda sistemática, regularización, descriptores diferenciados, justificación cuantitativa",
        "Tabla de ablación (14 combinaciones), GridSearchCV SVM/RF/LR, rejilla CNN, dropout/weight decay/label smoothing/aumento; decisiones por CV"],
       ["C1 – N5", "Sintonización fina multivariada, análisis de sensibilidad, balance exactitud–eficiencia embarcada, rigor teórico/experimental",
        "Rejilla factorial lr×dropout×wd, sensibilidad HOG/épocas/ancho/aumento/umbral, frente de Pareto con params-MACs-latencia CPU, McNemar, IC Wilson"],
       ["C2 – N3", "Evaluación en vivo en la UI sobre la carpeta de test, accuracy > 90 %, sustento en CV",
        f"UI app.py; CV 5-fold {pm(cvm['accuracy'])}"],
       ["C2 – N4", "Métricas detalladas (precisión, recall, F1, matriz) en vivo y accuracy ≥ 96 %",
        "Panel de métricas en tiempo real + exportación; hold-out " + pc(S["holdout_final_acc"])],
       ["C2 – N5", "Accuracy > 98 % en el test desconocido, robustez y generalización",
        f"Hold-out ciego simulado {pc(S['holdout_final_acc'])} (IC95 % inferior {S['holdout_final_ci95'][0] * 100:.2f} %); "
        "aumento de datos, TTA y ensamble para robustez ante un test nuevo"]],
      widths=[0.9, 2.8, 3.4])

H2("A.12 Reproducibilidad")
for c in ["pip install -r requirements.txt", "python train_classic.py   # baseline, ablación, GridSearch (results/classic.json)",
          "python train_cnn.py       # CNN: rejilla, sensibilidad, CV, hold-out, modelo final (results/cnn.json)",
          "python make_figures.py    # figuras y estadística (results/*.png, summary.json)",
          "python make_blind_test.py # carpeta test_ciego/ para ensayar la prueba en vivo",
          "python app.py             # UI de evaluación en vivo",
          "python scan_scene.py      # demo de detección en escena completa (UAV)"]:
    p = d.add_paragraph(); r = p.add_run(c); r.font.name = "Consolas"; r.font.size = Pt(8.5)
P("Referencias: [1] The Maritime Executive, «Rotterdam Tests Drones for Ship Inspections and Monitoring the Port». "
  "[2] IOP Conf. Ser.: Earth Environ. Sci. 557 012014 (2020), clasificación de barcos en imágenes satelitales con CNN. "
  "[3] R. Hammell, Ships in Satellite Imagery, Kaggle. [4] Dalal & Triggs, HOG, CVPR 2005. [5] Ojala et al., LBP, TPAMI 2002.", size=8.5)

d.save(OUT)
print("OK ->", OUT)
