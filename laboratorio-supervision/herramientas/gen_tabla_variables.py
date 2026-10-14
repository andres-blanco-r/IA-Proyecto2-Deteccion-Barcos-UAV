"""Genera plc/tablas/Variables_PLC.xlsx con el formato de importación de TIA Portal.

TIA Portal: Variables PLC > botón "Importar" (icono de flecha) > elegir el .xlsx.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).parent))
from variables_plc import VARIABLES  # noqa: E402

RAIZ = Path(__file__).resolve().parents[1]
SALIDA = RAIZ / "plc" / "tablas" / "Variables_PLC.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "PLC Tags"
ws.append(["Name", "Path", "Data Type", "Logical Address", "Comment", "Hmi Visible",
           "Hmi Accessible", "Hmi Writeable", "Typeobject ID", "Version ID"])
for c in ws[1]:
    c.font = Font(bold=True)
for nombre, tipo, dir_, com, tabla in VARIABLES:
    ws.append([nombre, f"Tabla_{tabla}", tipo, dir_, com, "True", "True", "False", "", ""])
for col, ancho in zip("ABCDE", (22, 18, 10, 16, 70)):
    ws.column_dimensions[col].width = ancho
wb.create_sheet("User Constants").append(["Name", "Path", "Data Type", "Value", "Comment"])
SALIDA.parent.mkdir(parents=True, exist_ok=True)
wb.save(SALIDA)

md = ["| Nombre | Tipo | Dirección | Tabla | Comentario |", "|---|---|---|---|---|"]
md += [f"| `{n}` | {t} | `{d}` | {tb} | {c} |" for n, t, d, c, tb in VARIABLES]
(RAIZ / "plc" / "tablas" / "Variables_PLC.md").write_text("\n".join(md) + "\n", encoding="utf-8")
print(SALIDA)
